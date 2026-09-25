"""Cell-by-cell comparison of two .xlsx files (values, number formats, fills, fonts, borders, widths,
   hidden rows/cols, merges, freeze panes, filter, validation/CF counts, notes, links, names).
   Usage: python compare_workbooks.py original.xlsx roundtripped.xlsx"""
import openpyxl, sys
def col(x):
  x=getattr(x,'rgb',x)
  if not isinstance(x,str): return None
  x=x[-6:].upper()
  return None if x=='000000' else x
a=openpyxl.load_workbook(sys.argv[1]); b=openpyxl.load_workbook(sys.argv[2])
diffs=[]
for ws in a:
  if ws.title not in b.sheetnames: diffs.append(('missing sheet',ws.title)); continue
  wb=b[ws.title]
  for row in ws.iter_rows():
    for c in row:
      d=wb[c.coordinate]
      if not ((isinstance(c.value,(int,float)) and not isinstance(c.value,bool) and isinstance(d.value,(int,float)) and not isinstance(d.value,bool) and abs(c.value-d.value)<=1e-9*max(1,abs(c.value))) or (c.value==d.value and type(c.value)==type(d.value))): diffs.append((c.coordinate,'value',repr(c.value),repr(d.value)))
      if c.number_format!=d.number_format: diffs.append((c.coordinate,'numfmt',c.number_format,d.number_format))
      if (col(c.fill.fgColor) if c.fill.fill_type else None)!=(col(d.fill.fgColor) if d.fill.fill_type else None): diffs.append((c.coordinate,'fill',c.fill.fgColor.rgb,d.fill.fgColor.rgb))
      fa,fb=c.font,d.font
      if (bool(fa.b),bool(fa.i),fa.u,fa.sz or 11,col(fa.color))!=(bool(fb.b),bool(fb.i),fb.u,fb.sz or 11,col(fb.color)): diffs.append((c.coordinate,'font',(fa.b,fa.i,fa.u,fa.sz,col(fa.color)),(fb.b,fb.i,fb.u,fb.sz,col(fb.color))))
      if c.alignment.horizontal!=d.alignment.horizontal or c.alignment.wrap_text!=d.alignment.wrap_text: diffs.append((c.coordinate,'align',c.alignment.horizontal,d.alignment.horizontal))
      for side in 'left right top bottom'.split():
        if (getattr(c.border,side) and getattr(c.border,side).style)!=(getattr(d.border,side) and getattr(d.border,side).style): diffs.append((c.coordinate,'border-'+side,getattr(c.border,side).style,getattr(d.border,side).style)); break
  for k,dim in ws.column_dimensions.items():
    if abs((dim.width or 0)-(wb.column_dimensions[k].width or 0))>0.6: diffs.append(('col '+k,'width',dim.width,wb.column_dimensions[k].width))
  if ws.freeze_panes!=wb.freeze_panes: diffs.append(('freeze',ws.freeze_panes,wb.freeze_panes))
  if str(ws.auto_filter.ref)!=str(wb.auto_filter.ref): diffs.append(('filter',ws.auto_filter.ref,wb.auto_filter.ref))
  if set(map(str,ws.merged_cells.ranges))!=set(map(str,wb.merged_cells.ranges)): diffs.append(('merges',))
  if set(ws.column_dimensions[k].hidden for k in ws.column_dimensions)!=set(wb.column_dimensions[k].hidden for k in ws.column_dimensions): diffs.append(('hidden cols',))
  for k,dim in ws.row_dimensions.items():
    if dim.hidden!=wb.row_dimensions[k].hidden: diffs.append(('row '+str(k),'hidden',dim.hidden,wb.row_dimensions[k].hidden))
    if dim.height and abs((dim.height or 0)-(wb.row_dimensions[k].height or 0))>0.6: diffs.append(('row '+str(k),'height',dim.height,wb.row_dimensions[k].height))
  if len(ws.data_validations.dataValidation)!=len(wb.data_validations.dataValidation): diffs.append(('dv count',len(ws.data_validations.dataValidation),len(wb.data_validations.dataValidation)))
  if len(ws.conditional_formatting)!=len(wb.conditional_formatting): diffs.append(('cf count',len(ws.conditional_formatting),len(wb.conditional_formatting)))
  for row in ws.iter_rows():
    for c in row:
      d=wb[c.coordinate]
      if (c.comment is None)!=(d.comment is None) or (c.comment and c.comment.text.strip() not in d.comment.text): diffs.append((c.coordinate,'comment',c.comment and c.comment.text, d.comment and d.comment.text))
      if (c.hyperlink and c.hyperlink.target)!=(d.hyperlink and d.hyperlink.target): diffs.append((c.coordinate,'link',c.hyperlink and c.hyperlink.target,d.hyperlink and d.hyperlink.target))
if set(a.defined_names.keys())!=set(b.defined_names.keys()): diffs.append(('defined names',list(a.defined_names.keys()),list(b.defined_names.keys())))
from collections import Counter
print(len(diffs),'diffs'); print(Counter(str(d[1]) if len(d)>1 else d[0] for d in diffs))
for d in diffs[:25]: print(d)
