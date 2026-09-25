"""Builds the two test workbooks used to compare Cellbook with real Excel.
   out/formulas.xlsx : 155 everyday formulas in column A (rows 10+) over a small data table
   out/stress.xlsx   : formatting, merges, notes, links, names, validation, CF, hidden rows/cols...
Run: python scripts/excel-oracle/make_workbooks.py   (needs: pip install openpyxl)"""
import os, datetime, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
OUT = os.path.join(os.path.dirname(__file__), 'out'); os.makedirs(OUT, exist_ok=True)

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'S'
for r in [['Code','Name','Qty','Rate','Date','Text','Blank'],
 ['A1','Alpha',10,2.5,datetime.date(2026,1,15),' pad ',None],
 ['B2','beta',20,3.125,datetime.date(2026,2,28),'x,y,z',None],
 ['C3','Gamma',0,-4.5,datetime.date(2024,2,29),'Hello World',None],
 ['D4','delta',35,1e-3,datetime.date(2026,12,31),'00123',None],
 ['E5','Alpha',15,7.75,datetime.date(2025,6,1),'',None]]: ws.append(r)
F = ['=SUM(C2:C6)','=AVERAGE(C2:C6)','=AVERAGEA(C2:C6)','=COUNT(A2:G6)','=COUNTA(A2:G6)','=COUNTBLANK(A2:G6)','=MIN(D2:D6)','=MAX(D2:D6)','=MEDIAN(C2:C6)','=MODE(C2:C6)','=STDEV(C2:C6)','=STDEVP(C2:C6)','=VAR(C2:C6)',
 '=ROUND(D3,2)','=ROUND(2.5,0)','=ROUND(-2.5,0)','=ROUND(1234.5678,-2)','=ROUNDUP(3.141,2)','=ROUNDDOWN(-3.149,2)','=MROUND(17,5)','=CEILING(4.2,1)','=FLOOR(4.8,1)','=INT(-3.5)','=TRUNC(-3.5)','=MOD(-7,3)','=MOD(7,-3)','=ABS(-4)','=POWER(2,10)','=SQRT(2)','=10/3','=0.1+0.2','=1E+15+1',
 '=SUMIF(B2:B6,"Alpha",C2:C6)','=SUMIF(B2:B6,"alpha",C2:C6)','=SUMIF(C2:C6,">=15")','=SUMIF(B2:B6,"*a",C2:C6)','=SUMIF(B2:B6,"?eta",C2:C6)','=COUNTIF(B2:B6,"<>Alpha")','=COUNTIF(C2:C6,0)','=COUNTIF(G2:G6,"")','=SUMIFS(C2:C6,B2:B6,"Alpha",C2:C6,">10")','=COUNTIFS(B2:B6,"*a*",D2:D6,">0")','=AVERAGEIF(B2:B6,"Alpha",C2:C6)',
 '=VLOOKUP("B2",A2:D6,2,FALSE)','=VLOOKUP("b2",A2:D6,4,0)','=VLOOKUP(12,C2:D6,2,TRUE)','=VLOOKUP("Z9",A2:D6,2,FALSE)','=HLOOKUP("Qty",A1:G6,3,FALSE)','=INDEX(A2:D6,3,2)','=MATCH("Gamma",B2:B6,0)','=MATCH(25,{10,20,30},1)','=_xlfn.XLOOKUP("D4",A2:A6,C2:C6)','=_xlfn.XLOOKUP("Q",A2:A6,C2:C6,"none")','=LOOKUP(2,{1,2,3},{"a","b","c"})','=CHOOSE(2,"x","y","z")',
 '=IF(C4=0,"zero","nz")','=IF(G2="","blank","no")','=IF(AND(C2>5,C3>5),1,0)','=IF(OR(C4>5,C5>50),1,0)','=NOT(TRUE)','=IFERROR(C2/C4,"div")','=_xlfn.IFNA(VLOOKUP("Z",A2:A6,1,0),"na")','=ISBLANK(G2)','=ISNUMBER(F5)','=ISTEXT(F5)','=ISERROR(1/0)','=_xlfn.IFS(C2>30,"a",C2>5,"b",TRUE,"c")','=_xlfn.SWITCH(2,1,"one",2,"two","other")',
 '=LEN(F2)','=TRIM(F2)','=UPPER(B3)','=LOWER(B2)','=PROPER("hello wORLD")','=LEFT(F4,5)','=RIGHT(F4,5)','=MID(F4,7,3)','=FIND("o",F4)','=SEARCH("W",F4)','=SUBSTITUTE(F3,",",";")','=REPLACE(F4,1,5,"Bye")','=REPT("-",3)','=CONCATENATE(A2,"-",B2)','=A2&"-"&C2','=TEXT(D3,"0.00")','=TEXT(1234567.891,"#,##0.00")','=TEXT(0.256,"0.0%")','=TEXT(E2,"dd-mm-yyyy")','=TEXT(E2,"mmm yyyy")','=TEXT(E2,"dddd")','=VALUE("12.5")','=VALUE(F5)','=F5+0','=EXACT("a","A")','=CHAR(65)','=CODE("A")','=CLEAN("a"&CHAR(10)&"b")','=T(B2)','=N(E2)',
 '=DATE(2026,2,30)','=YEAR(E3)','=MONTH(E3)','=DAY(E3)','=WEEKDAY(E2)','=WEEKDAY(E2,2)','=E3-E2','=E2+30','=EDATE(E2,1)','=EOMONTH(E2,0)','=EOMONTH(E4,1)','=DATEDIF(E4,E2,"y")','=DATEDIF(E4,E2,"m")','=DATEDIF(E4,E2,"d")','=NETWORKDAYS(E2,E3)','=WORKDAY(E2,10)','=_xlfn.DAYS(E3,E2)','=YEARFRAC(E2,E3,1)','=WEEKNUM(E5)','=TIME(14,30,0)','=HOUR(0.75)','=MINUTE(0.7513)','=DATEVALUE("2026-03-15")',
 '=SUMPRODUCT(C2:C6,D2:D6)','=SUMPRODUCT((B2:B6="Alpha")*C2:C6)','=PRODUCT(C2:C3)','=LARGE(C2:C6,2)','=SMALL(C2:C6,2)','=RANK(20,C2:C6)','=PERCENTILE(C2:C6,0.9)','=QUARTILE(C2:C6,1)','=COUNTA(B:B)','=SUM(C:C)','=MAX(E2:E6)-MIN(E2:E6)',
 '=PMT(0.01,12,10000)','=FV(0.05,10,-100)','=PV(0.05,10,-100)','=NPV(0.1,100,200,300)','=ROUND(IRR({-1000,300,400,500}),6)',
 '=SUM(C2:C6)/COUNT(C2:C6)','=C2*D2-C3/D3','=-C2^2','=2^3^2','=5%','=100*5%','="a"<"B"','=1="1"','=TRUE+1','=""=G2','=ROWS(A2:D6)','=COLUMNS(A2:D6)','=ROW(C4)','=COLUMN(C4)','=ADDRESS(2,3)','=INDIRECT("C"&2)','=OFFSET(A1,2,2)','=TRANSPOSE(C2)']
for i, f in enumerate(F): ws.cell(row=10 + i, column=1, value=f)
wb.save(os.path.join(OUT, 'formulas.xlsx'))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Dispatch Data'
ws.append(['Code','Name','Qty','Rate','Amount','Date','Pct','Flag','Note'])
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496'); c.alignment = Alignment(horizontal='center', wrap_text=True)
thin = Side(style='thin'); med = Side(style='medium', color='FF0000')
for i in range(2, 52):
  ws.append(['00%03d' % i, 'माल %d ट्रक' % i if i % 5 == 0 else 'Item %d' % i, i*3, 12.5+i/10, f'=C{i}*D{i}', datetime.date(2026,1,1)+datetime.timedelta(days=i), i/100, i % 2 == 0, None])
  for col, fmt in [('D','#,##0.00'),('E','_ * #,##0.00_ ;_ * -#,##0.00_ ;_ * "-"??_ ;_ @_ '),('F','dd-mm-yyyy'),('G','0.0%')]: ws[f'{col}{i}'].number_format = fmt
  for c in ws[i]: c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
ws['I5'] = '#N/A test'; ws['I6'] = '=1/0'; ws['I7'] = '=NA()'; ws['I8'] = "'quoted"; ws['I9'] = '="text"&" join"'
ws['K1'] = 'Total'; ws['K2'] = '=SUM(E2:E51)'; ws['K3'] = '=SUMIF(H2:H51,TRUE,E2:E51)'; ws['K4'] = '=Rates!$B$2*K2'; ws['K5'] = '=TaxRate*K2'; ws['K6'] = '=VLOOKUP("00010",A:E,5,FALSE)'
ws['K7'] = '=COUNTIF(B:B,"माल*")'; ws['K8'] = '=TEXT(F2,"dd mmm yyyy")'; ws['K9'] = '=ROUND(AVERAGE(D2:D51),2)'
ws['K2'].font = Font(bold=True, size=14, italic=True, underline='single', color='C00000'); ws['K2'].border = Border(bottom=med)
ws.merge_cells('M1:O2'); ws['M1'] = 'Merged header'; ws['M1'].alignment = Alignment(horizontal='center', vertical='center')
ws.freeze_panes = 'B2'; ws.column_dimensions['B'].width = 28; ws.column_dimensions['I'].hidden = True; ws.row_dimensions[10].hidden = True; ws.row_dimensions[3].height = 30
ws['B3'].alignment = Alignment(wrap_text=True); ws['B3'] = 'Long text that wraps across lines in this cell for testing'
ws['A1'].comment = Comment('Code has leading zeros', 'Dhruv')
ws['B4'].hyperlink = 'https://example.com'; ws['B4'].style = 'Hyperlink'
dv = DataValidation(type='list', formula1='"Yes,No,Maybe"', allow_blank=True); ws.add_data_validation(dv); dv.add('Q2:Q20')
ws.conditional_formatting.add('C2:C51', CellIsRule(operator='greaterThan', formula=['100'], fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006')))
r = wb.create_sheet('Rates'); r['A1'] = 'Key'; r['B1'] = 'Value'; r['A2'] = 'Freight'; r['B2'] = 1.18; r['C2'] = '=B2*100'
wb.defined_names['TaxRate'] = DefinedName('TaxRate', attr_text='Rates!$B$2')
wb.save(os.path.join(OUT, 'stress.xlsx'))
print('wrote', OUT)
